"""
เรนเดอร์ LayoutSpec เป็นไฟล์ .pptx ด้วย python-pptx

จุดที่ต้องระวังที่สุดคือฟอนต์ไทย: run.font.name ของ python-pptx เขียนแค่ <a:latin>
ซึ่ง PowerPoint ไม่ได้ใช้กับอักษรไทย ต้องเขียน <a:cs> (complex script) ให้ด้วย
ไม่งั้นข้อความไทยจะถูกเรนเดอร์ด้วยฟอนต์ดีฟอลต์แทน
"""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

import theme
from spec import Arrow, Badge, Box, Cylinder, Person, Text, expand_paras, segments

EMU_PER_INCH = 914400

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCHOR = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}


def E(inches: float) -> Emu:
    return Emu(int(round(inches * EMU_PER_INCH)))


def C(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def set_run_font(run, name: str, bold: bool, size_pt: float, color: str):
    """
    ตั้งฟอนต์ให้ครบทั้งสามชุด: latin / complex script / east asian
    <a:cs> คือชุดที่ PowerPoint ใช้กับอักษรไทย
    """
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = C(color)
    run.font.name = name          # เขียน <a:latin> ให้
    rPr = run._r.get_or_add_rPr()
    # ลำดับของลูกใน <a:rPr> ถูกกำหนดไว้ตายตัวว่า latin → ea → cs
    # ถ้าสลับลำดับ PowerPoint จะมองว่าไฟล์เสียและขอซ่อมก่อนเปิด
    anchor = rPr.find(qn("a:latin"))
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            anchor.addnext(el)
        el.set("typeface", name)
        anchor = el


def no_shadow(shape):
    shape.shadow.inherit = False


def add_box(slide, b: Box):
    if b.radius > 0:
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, E(b.x), E(b.y), E(b.w), E(b.h)
        )
        shape.adjustments[0] = min(0.5, b.radius / min(b.w, b.h))
    else:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(b.x), E(b.y), E(b.w), E(b.h))

    if b.fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = C(b.fill)
    else:
        shape.fill.background()

    if b.line:
        shape.line.color.rgb = C(b.line)
        shape.line.width = Pt(b.line_w)
    else:
        shape.line.fill.background()

    no_shadow(shape)
    shape.text_frame.word_wrap = False
    return shape


def add_text(slide, t: Text):
    tb = slide.shapes.add_textbox(E(t.x), E(t.y), E(t.w), E(t.h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCHOR[t.valign]

    paras = expand_paras(t.paras)
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = ALIGN[p.align]
        para.line_spacing = t.spacing
        para.space_before = Pt(p.space_before) if i > 0 and p.space_before else Pt(0)
        para.space_after = Pt(0)
        _set_bullet(para, p)
        # แยก run ตามฟอนต์ เพื่อให้ลูกศรใช้ฟอนต์สำรองที่มี glyph จริง
        for chunk, font_name in segments(p.text, p.font):
            run = para.add_run()
            run.text = chunk
            set_run_font(run, font_name, p.bold, p.size, p.color)
    return tb


def _set_bullet(para, p):
    """
    จุดนำแบบย่อหน้าแขวน — ให้ PowerPoint เป็นคนวาดจุดนำและเยื้องบรรทัดต่อเอง
    ถ้าใส่ '•' ลงไปในข้อความตรง ๆ พอตัดบรรทัด จุดนำจะหลุดไปอยู่บรรทัดเดียวลอย ๆ
    """
    pPr = para._pPr if para._pPr is not None else para._p.get_or_add_pPr()
    if not p.bullet and not p.auto_num:
        pPr.append(pPr.makeelement(qn("a:buNone"), {}))
        return

    marl = int(round(p.indent * EMU_PER_INCH))
    pPr.set("marL", str(marl))
    pPr.set("indent", str(-marl))
    font_el = pPr.makeelement(qn("a:buFont"), {"typeface": theme.FONT_TH})
    pPr.append(font_el)
    if p.auto_num:
        pPr.append(pPr.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"}))
    else:
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": p.bullet}))


def add_arrow(slide, a: Arrow):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(a.x1), E(a.y1), E(a.x2), E(a.y2))
    conn.line.color.rgb = C(a.color)
    conn.line.width = Pt(a.width)
    ln = conn.line._get_or_add_ln()
    tail = ln.find(qn("a:tailEnd"))
    if tail is None:
        tail = ln.makeelement(qn("a:tailEnd"), {})
        ln.append(tail)
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return conn


def add_badge(slide, b: Badge):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, E(b.cx - b.r), E(b.cy - b.r), E(2 * b.r), E(2 * b.r)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = C(b.fill)
    shape.line.fill.background()
    no_shadow(shape)
    tf = shape.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = tf.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    para.line_spacing = 1.0
    run = para.add_run()
    run.text = b.text
    set_run_font(run, theme.FONT_TH, True, b.size, b.color)
    return shape


def add_cylinder(slide, c: Cylinder):
    shape = slide.shapes.add_shape(MSO_SHAPE.CAN, E(c.x), E(c.y), E(c.w), E(c.h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = C(c.fill)
    shape.line.color.rgb = C(c.line)
    shape.line.width = Pt(1.25)
    no_shadow(shape)
    shape.text_frame.word_wrap = False
    return shape


def add_person(slide, p: Person):
    head_d = p.h * 0.42
    head = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, E(p.cx - head_d / 2), E(p.top), E(head_d), E(head_d)
    )
    body_w = p.h * 0.80
    body_top = p.top + head_d + p.h * 0.10
    body_h = (p.top + p.h + p.h * 0.18) - body_top
    body = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, E(p.cx - body_w / 2), E(body_top), E(body_w), E(body_h)
    )
    body.adjustments[0] = min(0.5, (p.h * 0.22) / min(body_w, body_h))
    for shape in (head, body):
        shape.fill.solid()
        shape.fill.fore_color.rgb = C(p.color)
        shape.line.fill.background()
        no_shadow(shape)
        shape.text_frame.word_wrap = False
    return head, body


DISPATCH = {
    Box: add_box,
    Text: add_text,
    Arrow: add_arrow,
    Badge: add_badge,
    Cylinder: add_cylinder,
    Person: add_person,
}


def render_deck(deck, out_path):
    prs = Presentation()
    prs.slide_width = E(theme.SLIDE_W)
    prs.slide_height = E(theme.SLIDE_H)
    blank = prs.slide_layouts[6]

    for s in deck.slides:
        slide = prs.slides.add_slide(blank)
        _set_background(slide, theme.PAGE)
        for item in s.primitives:
            DISPATCH[type(item)](slide, item)
        if s.notes:
            slide.notes_slide.notes_text_frame.text = s.notes

    prs.save(str(out_path))
    return out_path


def _set_background(slide, hex_color: str):
    """พื้นหลังสไลด์ทึบ — ไม่ปล่อยให้ไปพึ่งค่าดีฟอลต์ของธีม"""
    bg = slide.background
    bg.fill.solid()
    bg.fill.fore_color.rgb = C(hex_color)


def slide_text_dump(path) -> list[list[str]]:
    """อ่านข้อความทุกสไลด์กลับออกมา ใช้แทน markitdown ที่เครื่องนี้ไม่มี"""
    prs = Presentation(str(path))
    out = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
        out.append(texts)
    return out
