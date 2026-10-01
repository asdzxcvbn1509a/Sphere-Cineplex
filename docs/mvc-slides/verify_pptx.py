"""
ตรวจไฟล์ .pptx ที่สร้างเสร็จแล้ว — แทน markitdown และ validate.py ที่เครื่องนี้ไม่มี

    python docs/mvc-slides/verify_pptx.py
"""

import re
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

from pptx import Presentation  # noqa: E402

from endpoints import ENDPOINTS  # noqa: E402

DECK = HERE.parent.parent / "MVC-Architecture-API.pptx"
THAI = re.compile(r"[฀-๿]")
PLACEHOLDER = re.compile(r"\bTODO\b|lorem|ipsum|\[insert|\bx{3,}\b", re.IGNORECASE)


A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"

# ลำดับลูกที่สเปก DrawingML กำหนดไว้ตายตัว — สลับลำดับแล้ว PowerPoint จะขอซ่อมไฟล์ก่อนเปิด
CHILD_ORDER = {
    f"{A}rPr": [
        "ln", "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill",
        "effectLst", "effectDag", "highlight", "uLnTx", "uLn", "uFillTx", "uFill",
        "latin", "ea", "cs", "sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst",
    ],
    f"{A}pPr": [
        "lnSpc", "spcBef", "spcAft", "buClrTx", "buClr", "buSzTx", "buSzPct", "buSzPts",
        "buFontTx", "buFont", "buNone", "buAutoNum", "buChar", "tabLst", "defRPr", "extLst",
    ],
    f"{A}ln": [
        "noFill", "solidFill", "gradFill", "pattFill", "prstDash", "custDash",
        "round", "bevel", "miter", "headEnd", "tailEnd", "extLst",
    ],
}


def check_element_order(path) -> list[str]:
    from lxml import etree

    errors = []
    with zipfile.ZipFile(path) as zf:
        for name in sorted(n for n in zf.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)):
            root = etree.fromstring(zf.read(name))
            for tag, order in CHILD_ORDER.items():
                for parent in root.iter(tag):
                    seen = [c.tag.split("}")[-1] for c in parent if isinstance(c.tag, str)]
                    ranks = [order.index(c) for c in seen if c in order]
                    if ranks != sorted(ranks):
                        errors.append(f"{name} <{tag.split('}')[-1]}> ลำดับลูก {seen}")
    return errors


def main() -> int:
    problems: list[str] = []
    prs = Presentation(str(DECK))
    slides = list(prs.slides)
    print(f"ไฟล์            : {DECK.name}  ({DECK.stat().st_size / 1024:.0f} KB)")
    print(f"จำนวนสไลด์       : {len(slides)}")
    print(f"ขนาดสไลด์        : {prs.slide_width / 914400:.3f}\" x {prs.slide_height / 914400:.3f}\"")

    expected = 3 + 8 + len(ENDPOINTS)
    if len(slides) != expected:
        problems.append(f"จำนวนสไลด์ {len(slides)} ไม่เท่ากับที่คาดไว้ {expected}")

    # ---- ข้อความครบทุก endpoint หรือไม่ ----
    all_text = []
    for slide in slides:
        texts = [
            shape.text_frame.text
            for shape in slide.shapes
            if shape.has_text_frame and shape.text_frame.text.strip()
        ]
        all_text.append("\n".join(texts))
    joined = "\n".join(all_text)

    missing = [f"{ep.method} {ep.path}" for ep in ENDPOINTS if f"{ep.method}  {ep.path}" not in joined]
    print(f"endpoint ที่พบ   : {len(ENDPOINTS) - len(missing)} / {len(ENDPOINTS)}")
    if missing:
        problems.append("ไม่พบ endpoint บนสไลด์: " + ", ".join(missing[:5]))

    titles_missing = [ep.title for ep in ENDPOINTS if ep.title not in joined]
    if titles_missing:
        problems.append("ไม่พบหัวข้อสไลด์: " + ", ".join(titles_missing[:5]))

    for i, text in enumerate(all_text, start=1):
        hit = PLACEHOLDER.search(text)
        if hit:
            problems.append(f"หน้า {i}: พบข้อความค้างไว้ '{hit.group(0)}'")

    # ---- ฟอนต์ไทยถูกตั้งเป็น complex script ทุก run หรือไม่ ----
    bad_font = []
    with zipfile.ZipFile(DECK) as zf:
        slide_parts = sorted(n for n in zf.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n))
        for name in slide_parts:
            xml = zf.read(name).decode("utf-8")
            for run in re.findall(r"<a:r>.*?</a:r>", xml, flags=re.S):
                text = "".join(re.findall(r"<a:t>(.*?)</a:t>", run, flags=re.S))
                if THAI.search(text) and 'a:cs typeface="Leelawadee UI"' not in run:
                    bad_font.append((name, text[:30]))
    print(f"ส่วน XML ของสไลด์ : {len(slide_parts)} ไฟล์")
    print(f"run ภาษาไทยที่ไม่ได้ตั้งฟอนต์ complex script : {len(bad_font)}")
    if bad_font:
        problems.append(f"พบ run ภาษาไทยที่ไม่ได้ตั้ง <a:cs> {len(bad_font)} จุด เช่น {bad_font[0]}")

    # ---- ลำดับลูกของ element ที่เราเขียน XML เอง ----
    order_errors = check_element_order(DECK)
    print(f"element ที่ลำดับลูกผิดสเปก : {len(order_errors)}")
    if order_errors:
        problems.append(f"ลำดับลูกผิดสเปก {len(order_errors)} จุด เช่น {order_errors[0]}")

    # ---- จุดนำแบบย่อหน้าแขวนถูกเขียนลงไปจริงหรือไม่ ----
    with zipfile.ZipFile(DECK) as zf:
        sample = zf.read(slide_parts[6]).decode("utf-8")
    for tag, label in (("a:buChar", "จุดนำ"), ("a:buAutoNum", "รายการมีหมายเลข"), ("indent=", "ย่อหน้าแขวน")):
        if tag not in sample:
            problems.append(f"ไม่พบ {label} ({tag}) ในสไลด์ตัวอย่าง")

    print()
    if problems:
        print(f"ตรวจไฟล์: พบ {len(problems)} รายการ")
        for p in problems:
            print("  -", p)
        return 1
    print("ตรวจไฟล์: ผ่านทุกข้อ")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
